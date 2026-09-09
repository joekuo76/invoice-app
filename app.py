import io
import os
import re
from datetime import datetime

import numpy as np
import openpyxl
import pandas as pd
import streamlit as st
from PIL import Image, ImageOps
import zxingcpp

st.set_page_config(page_title="請款單自動生成系統", page_icon="🧾", layout="wide")
st.title("🧾 請款單自動生成工具 V5")
st.caption("台灣電子發票 QR 自動讀取｜收據/停車單快速人工輸入｜一鍵產生 Excel")
st.success("輕量穩定版：不使用 Gemini、不使用 EasyOCR、不需要 API Key。")

TEMPLATE_PATH = "template.xlsx"
MAX_ITEMS = 8

if "items_v5" not in st.session_state:
    st.session_state.items_v5 = []

with st.sidebar:
    st.header("⚙️ 請款資料")
    applicant = st.text_input("請款人姓名")
    department = st.text_input("請款單位")
    st.caption("資料確認後會自動帶入 Excel。")


def load_image(data: bytes):
    img = Image.open(io.BytesIO(data))
    return ImageOps.exif_transpose(img).convert("RGB")


def roc_date_to_iso(v: str):
    digits = re.sub(r"\D", "", v or "")
    if len(digits) != 7:
        return ""
    try:
        return datetime(
            int(digits[:3]) + 1911,
            int(digits[3:5]),
            int(digits[5:7]),
        ).strftime("%Y-%m-%d")
    except Exception:
        return ""


def decode_qr(img):
    try:
        results = zxingcpp.read_barcodes(np.array(img))
        return [r.text.strip() for r in results if getattr(r, "text", "").strip()]
    except Exception:
        return []


def parse_invoice_qr(text):
    # 台灣電子發票左側 QR 的前段：
    # 10碼發票號碼 + 7碼民國日期 + 4碼隨機碼 + 8碼銷售額 + 8碼總額
    s = (text or "").replace("\n", "").strip()
    if len(s) < 37:
        return None

    invoice_no = s[:10].upper()
    roc_date = s[10:17]
    total_hex = s[29:37]

    if not re.fullmatch(r"[A-Z]{2}\d{8}", invoice_no):
        return None
    if not re.fullmatch(r"\d{7}", roc_date):
        return None
    try:
        amount = int(total_hex, 16)
    except Exception:
        return None

    return {
        "document_type": "發票",
        "date": roc_date_to_iso(roc_date),
        "reference_no": invoice_no,
        "summary": "",
        "amount": amount,
        "payment_method": "未知",
        "location": "",
        "plate_no": "",
        "source": "QR自動",
    }


def scan_invoice(data):
    img = load_image(data)
    qr_texts = decode_qr(img)
    for text in qr_texts:
        parsed = parse_invoice_qr(text)
        if parsed:
            return parsed
    return None


def normalize_item(x):
    x = dict(x)
    x["document_type"] = str(x.get("document_type", "其他") or "其他")
    x["date"] = str(x.get("date", "") or "").strip()
    x["reference_no"] = str(x.get("reference_no", "") or "").upper().strip()
    x["summary"] = str(x.get("summary", "") or "").strip()
    x["payment_method"] = str(x.get("payment_method", "未知") or "未知")
    x["location"] = str(x.get("location", "") or "").strip()
    x["plate_no"] = str(x.get("plate_no", "") or "").upper().strip()
    try:
        x["amount"] = int(float(x.get("amount", 0) or 0))
    except Exception:
        x["amount"] = 0
    return x


def write_excel(items, dept, applicant_name):
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError("找不到 template.xlsx")

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb.active

    # 依目前 template.xlsx 的固定版型：
    # 明細 5~12、合計 13、簽名內容 16
    start_row, end_row, total_row = 5, 12, 13

    if len(items) > MAX_ITEMS:
        raise ValueError(f"目前範本最多 {MAX_ITEMS} 筆")

    ws["B3"] = dept or ""
    ws["F3"] = datetime.now().strftime("%Y-%m-%d")

    for r in range(start_row, end_row + 1):
        for c in range(1, 8):
            ws.cell(r, c).value = None

    for idx, raw in enumerate(items):
        item = normalize_item(raw)
        r = start_row + idx

        summary = item["summary"]
        if item["document_type"] == "停車繳費單" and not summary:
            summary = "停車費"
        if item["document_type"] == "停車繳費單" and item["location"]:
            summary = f"{summary or '停車費'}－{item['location']}"

        notes = []
        if item["plate_no"]:
            notes.append(f"車牌:{item['plate_no']}")
        if item["location"] and item["document_type"] != "停車繳費單":
            notes.append(item["location"])

        ws.cell(r, 1, item["date"])
        ws.cell(r, 2, item["reference_no"])
        ws.cell(r, 3, summary)
        ws.cell(r, 4, item["amount"]).number_format = '$#,##0'
        ws.cell(r, 5, item["payment_method"])
        ws.cell(r, 6, item["document_type"])
        ws.cell(r, 7, " / ".join(notes))

    ws.cell(total_row, 4, f"=SUM(D{start_row}:D{end_row})").number_format = '$#,##0'

    if applicant_name:
        ws["G16"] = applicant_name

    # 修正列印只印部分範圍的問題
    ws.print_area = "A1:G16"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.print_options.horizontalCentered = True

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


tab1, tab2 = st.tabs(["📷 電子發票 QR", "✍️ 收據 / 停車單"])

with tab1:
    st.subheader("電子發票")
    st.write("上傳或用手機拍攝有 QR Code 的台灣電子發票。")

    invoice_files = st.file_uploader(
        "選擇電子發票照片（可多選）",
        type=["jpg", "jpeg", "png", "webp"],
        accept_multiple_files=True,
        key="invoice_files",
    )

    if invoice_files:
        cols = st.columns(min(3, len(invoice_files)))
        for i, f in enumerate(invoice_files):
            with cols[i % len(cols)]:
                st.image(f, caption=f.name, width="stretch")

        if st.button("🔎 讀取發票 QR", type="primary", width="stretch"):
            added = 0
            for f in invoice_files:
                try:
                    result = scan_invoice(f.getvalue())
                    if result:
                        result["summary"] = st.session_state.get("default_invoice_summary", "")
                        st.session_state.items_v5.append(result)
                        added += 1
                    else:
                        st.warning(f"{f.name}：沒有讀到可解析的台灣電子發票 QR，請改用人工輸入。")
                except Exception as exc:
                    st.warning(f"{f.name}：讀取失敗（{exc}），請改用人工輸入。")
            if added:
                st.success(f"已加入 {added} 筆發票。")
                st.rerun()

    st.text_input(
        "發票預設摘要（選填，例如：加油、餐費、工程耗材）",
        key="default_invoice_summary",
    )

with tab2:
    st.subheader("收據 / 停車繳費單")
    st.write("特殊單據不跑大型 OCR，直接看照片快速輸入，穩定且沒有 API 額度。")

    manual_photo = st.file_uploader(
        "拍照 / 上傳單據（選填，方便對照）",
        type=["jpg", "jpeg", "png", "webp"],
        key="manual_photo",
    )

    left, right = st.columns([1, 1])

    with left:
        if manual_photo:
            st.image(manual_photo, caption=manual_photo.name, width="stretch")
        else:
            st.info("手機可直接從這裡拍照或選擇相簿照片。")

    with right:
        with st.form("manual_form", clear_on_submit=True):
            doc_type = st.selectbox("單據類型", ["停車繳費單", "收據", "其他"])
            date = st.text_input("日期", placeholder="2026-09-09")
            reference_no = st.text_input("單據 / 停車繳費編號")
            summary = st.text_input("付款內容摘要", value="停車費")
            amount = st.number_input("金額", min_value=0, step=1)
            payment_method = st.selectbox("付款方式", ["未知", "現金", "刷卡"])
            location = st.text_input("店家 / 停車場")
            plate_no = st.text_input("車牌號碼")
            submitted = st.form_submit_button("➕ 加入請款明細", type="primary", width="stretch")

            if submitted:
                st.session_state.items_v5.append({
                    "document_type": doc_type,
                    "date": date,
                    "reference_no": reference_no,
                    "summary": summary,
                    "amount": int(amount),
                    "payment_method": payment_method,
                    "location": location,
                    "plate_no": plate_no,
                    "source": "人工輸入",
                })
                st.rerun()


st.divider()
st.subheader("📋 請款明細確認")

if not st.session_state.items_v5:
    st.info("目前還沒有資料。請先讀取電子發票 QR，或新增收據 / 停車單。")
else:
    df = pd.DataFrame(st.session_state.items_v5)

    expected = [
        "document_type", "date", "reference_no", "summary", "amount",
        "payment_method", "location", "plate_no", "source",
    ]
    for col in expected:
        if col not in df.columns:
            df[col] = ""
    df = df[expected]

    edited = st.data_editor(
        df,
        hide_index=True,
        width="stretch",
        num_rows="dynamic",
        column_config={
            "document_type": st.column_config.SelectboxColumn(
                "單據類型", options=["發票", "停車繳費單", "收據", "其他"]
            ),
            "date": st.column_config.TextColumn("日期"),
            "reference_no": st.column_config.TextColumn("發票 / 單據編號"),
            "summary": st.column_config.TextColumn("付款摘要"),
            "amount": st.column_config.NumberColumn("金額", min_value=0, step=1, format="%d"),
            "payment_method": st.column_config.SelectboxColumn(
                "付款方式", options=["未知", "現金", "刷卡"]
            ),
            "location": st.column_config.TextColumn("店家 / 停車場"),
            "plate_no": st.column_config.TextColumn("車牌"),
            "source": st.column_config.TextColumn("來源", disabled=True),
        },
        key="items_editor",
    )

    st.session_state.items_v5 = edited.to_dict("records")

    total = sum(int(float(x.get("amount", 0) or 0)) for x in st.session_state.items_v5)
    c1, c2 = st.columns(2)
    c1.metric("目前筆數", len(st.session_state.items_v5))
    c2.metric("請款合計", f"${total:,}")

    if len(st.session_state.items_v5) > MAX_ITEMS:
        st.error(f"目前 Excel 範本最多 {MAX_ITEMS} 筆，請刪除部分資料或分成兩張請款單。")
    else:
        try:
            excel = write_excel(st.session_state.items_v5, department, applicant)
            st.download_button(
                "📥 下載請款單 Excel",
                data=excel,
                file_name=f"請款單_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                width="stretch",
            )
        except Exception as exc:
            st.error(f"Excel 產生失敗：{exc}")

    if st.button("🗑️ 清空本次請款資料", width="stretch"):
        st.session_state.items_v5 = []
        st.rerun()

with st.expander("ℹ️ V5 使用方式"):
    st.markdown("""
**電子發票：** 拍照或上傳 → 讀 QR → 自動取得發票號碼、日期與總金額 → 在明細表確認。

**停車單 / 一般收據：** 拍照放在左邊對照 → 右邊快速輸入日期、金額、單號、車牌等 → 加入明細。

**最後：** 檢查明細 → 下載 Excel。列印設定已固定為 A4 橫向並縮放至單頁。
""")
